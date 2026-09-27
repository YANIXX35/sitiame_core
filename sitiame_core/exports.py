# Copyright (c) 2026, Sitiame Capital
# License: MIT

"""List-view "Telecharger" button (public/js/list_download.js): the
documents of a list -- the checked rows, else every row matching the
list's current filters -- as PDF, Word, Excel or CSV.

Documents only ever come from frappe.get_list, i.e. with the caller's read
permission and User Permissions applied: a PME gets its own company's
documents and nothing else, without opening ERPNext's native "Export"
right (which would expose far more than these lists).
"""

import csv
import datetime
import decimal
import io

import frappe
from frappe import _

FORMATS = ("pdf", "docx", "xlsx", "csv")
# Laid-out formats render every document one by one: slow, so capped lower.
LIMITS = {"pdf": 200, "docx": 200, "xlsx": 5000, "csv": 5000}
BOOLEAN_FIELDS = {"is_stock_item"}

_LINES = [("item_name", "Article"), ("qty", "Qté"), ("rate", "Prix unitaire"), ("amount", "Montant")]
_STOCK_LINES = [("item_code", "Code"), ("item_name", "Article"), ("qty", "Qté"), ("rate", "Prix"), ("amount", "Montant"), ("warehouse", "Dépôt")]

EXPORTS = {
	"Sales Invoice": {
		"title": "Facture de vente",
		"file": "factures_de_vente",
		"columns": [
			("name", "N°"), ("posting_date", "Date"), ("customer_name", "Client"), ("net_total", "HT"),
			("total_taxes_and_charges", "TVA"), ("grand_total", "TTC"), ("outstanding_amount", "Reste à payer"), ("status", "Statut"),
		],
		"items": ("Sales Invoice Item", _LINES + [("income_account", "Compte")]),
	},
	"Purchase Invoice": {
		"title": "Facture d'achat",
		"file": "factures_d_achat",
		"columns": [
			("name", "N°"), ("posting_date", "Date"), ("supplier_name", "Fournisseur"), ("bill_no", "N° fournisseur"),
			("net_total", "HT"), ("total_taxes_and_charges", "TVA"), ("grand_total", "TTC"),
			("outstanding_amount", "Reste à payer"), ("status", "Statut"),
		],
		"items": ("Purchase Invoice Item", _LINES + [("expense_account", "Compte")]),
	},
	"Payment Entry": {
		"title": "Paiement",
		"file": "paiements",
		"columns": [
			("name", "N°"), ("posting_date", "Date"), ("payment_type", "Type"), ("party_name", "Tiers"),
			("mode_of_payment", "Moyen de paiement"), ("paid_amount", "Montant"), ("reference_no", "Référence"), ("status", "Statut"),
		],
		"items": (
			"Payment Entry Reference",
			[("reference_doctype", "Type de document"), ("reference_name", "Document réglé"), ("allocated_amount", "Montant alloué")],
		),
	},
	"Journal Entry": {
		"title": "Écriture de journal",
		"file": "ecritures_de_journal",
		"columns": [
			("name", "N°"), ("posting_date", "Date"), ("voucher_type", "Type"), ("user_remark", "Libellé"),
			("total_debit", "Total débit"), ("total_credit", "Total crédit"),
		],
		"items": (
			"Journal Entry Account",
			[("account", "Compte"), ("party", "Tiers"), ("debit_in_account_currency", "Débit"), ("credit_in_account_currency", "Crédit")],
		),
	},
	"Stock Entry": {
		"title": "Mouvement de stock",
		"file": "mouvements_de_stock",
		"columns": [
			("name", "N°"), ("posting_date", "Date"), ("stock_entry_type", "Type"),
			("total_outgoing_value", "Valeur sortie"), ("total_incoming_value", "Valeur entrée"),
		],
		"items": (
			"Stock Entry Detail",
			[
				("item_code", "Code"), ("item_name", "Article"), ("qty", "Qté"), ("s_warehouse", "Dépôt source"),
				("t_warehouse", "Dépôt cible"), ("basic_rate", "Prix"), ("amount", "Montant"),
			],
		),
	},
	"Purchase Receipt": {
		"title": "Réception d'achat",
		"file": "receptions_d_achat",
		"columns": [("name", "N°"), ("posting_date", "Date"), ("supplier_name", "Fournisseur"), ("grand_total", "Total"), ("status", "Statut")],
		"items": ("Purchase Receipt Item", _STOCK_LINES),
	},
	"Delivery Note": {
		"title": "Bon de livraison",
		"file": "bons_de_livraison",
		"columns": [("name", "N°"), ("posting_date", "Date"), ("customer_name", "Client"), ("grand_total", "Total"), ("status", "Statut")],
		"items": ("Delivery Note Item", _STOCK_LINES),
	},
	"Purchase Order": {
		"title": "Commande d'achat",
		"file": "commandes_d_achat",
		"columns": [("name", "N°"), ("transaction_date", "Date"), ("supplier_name", "Fournisseur"), ("grand_total", "Total"), ("status", "Statut")],
		"items": ("Purchase Order Item", _LINES + [("schedule_date", "Date prévue")]),
	},
	"Sales Order": {
		"title": "Commande client",
		"file": "commandes_clients",
		"columns": [
			("name", "N°"), ("transaction_date", "Date"), ("customer_name", "Client"), ("delivery_date", "Livraison"),
			("grand_total", "Total"), ("status", "Statut"),
		],
		"items": ("Sales Order Item", _LINES + [("delivery_date", "Date prévue")]),
	},
	"Item": {
		"title": "Article",
		"file": "articles_et_stock",
		"columns": [
			("item_code", "Code"), ("item_name", "Désignation"), ("item_group", "Groupe"), ("stock_uom", "Unité"),
			("is_stock_item", "Stocké"), ("valuation_rate", "Valorisation"), ("stock", "Stock actuel"),
		],
		"stock": True,
	},
}


def _config(doctype, fmt):
	if doctype not in EXPORTS:
		frappe.throw(_("Telechargement non disponible pour {0}.").format(doctype))
	if fmt not in FORMATS:
		frappe.throw(_("Format non pris en charge : {0}").format(fmt))
	return EXPORTS[doctype]


def _parse(value):
	return frappe.parse_json(value) if isinstance(value, str) and value else value


def _names(doctype, names, filters, limit):
	"""Names of the checked rows, else of the rows matching the list
	filters -- read-permission and User Permission checked, at most
	limit + 1 so "too many" can be told apart from "exactly the limit"."""
	names = _parse(names) or []
	return frappe.get_list(
		doctype,
		filters=[["name", "in", names]] if names else (_parse(filters) or []),
		pluck="name",
		order_by="creation asc",
		limit_page_length=limit + 1,
	)


@frappe.whitelist()
def count_documents(doctype, fmt, names=None, filters=None):
	_config(doctype, fmt)
	return {"count": len(_names(doctype, names, filters, LIMITS[fmt])), "limit": LIMITS[fmt]}


@frappe.whitelist()
def download(doctype, fmt, names=None, filters=None):
	config = _config(doctype, fmt)
	found = _names(doctype, names, filters, LIMITS[fmt])
	if not found:
		frappe.throw(_("Aucun document a telecharger."))
	if len(found) > LIMITS[fmt]:
		frappe.throw(
			_("Plus de {0} documents : filtrez davantage (par mois par exemple) pour ce format.").format(LIMITS[fmt])
		)

	if fmt == "pdf":
		content = build_pdf(doctype, found)
	else:
		builder = {"docx": build_docx, "xlsx": build_xlsx, "csv": build_csv}[fmt]
		content = builder(config, _rows(doctype, config, found))

	frappe.local.response.filename = f"{config['file']}_{frappe.utils.today()}.{fmt}"
	frappe.local.response.filecontent = content
	frappe.local.response.type = "download"


def _rows(doctype, config, names):
	fields = ["name"] + [field for field, label in config["columns"] if field not in ("name", "stock")]
	if doctype != "Item":
		fields.append("company")
	docs = frappe.get_list(
		doctype, filters=[["name", "in", names]], fields=fields, order_by="creation asc", limit_page_length=len(names)
	)

	if config.get("items"):
		child, columns = config["items"]
		lines = {}
		# parents already permission-filtered above
		for row in frappe.get_all(
			child,
			filters={"parent": ["in", names], "parenttype": doctype},
			fields=["parent"] + [field for field, label in columns],
			order_by="parent asc, idx asc",
		):
			lines.setdefault(row.parent, []).append(row)
		for doc in docs:
			doc["items"] = lines.get(doc.name, [])

	if config.get("stock"):
		warehouses = frappe.get_list("Warehouse", filters={"is_group": 0}, pluck="name", limit_page_length=0)
		stock = {}
		for row in frappe.get_all(
			"Bin",
			filters={"item_code": ["in", names], "warehouse": ["in", warehouses or [""]]},
			fields=["item_code", "actual_qty"],
		):
			stock[row.item_code] = stock.get(row.item_code, 0.0) + float(row.actual_qty or 0)
		for doc in docs:
			doc["stock"] = stock.get(doc.name, 0.0)

	return docs


def _number(value):
	text = f"{float(value):,.2f}".replace(",", " ").replace(".", ",")
	return text[:-3] if text.endswith(",00") else text


def _text(value, field=None):
	if field in BOOLEAN_FIELDS:
		return _("Oui") if value else _("Non")
	if value is None:
		return ""
	if isinstance(value, datetime.date):
		return value.strftime("%d/%m/%Y")
	if isinstance(value, (int, float, decimal.Decimal)) and not isinstance(value, bool):
		return _number(value)
	return str(value)


def _cell(value, field=None):
	if field in BOOLEAN_FIELDS:
		return _text(value, field)
	if isinstance(value, decimal.Decimal):
		return float(value)
	return value


def build_csv(config, docs):
	out = io.StringIO()
	writer = csv.writer(out, delimiter=";")
	writer.writerow([label for field, label in config["columns"]])
	for doc in docs:
		writer.writerow([_text(doc.get(field), field) for field, label in config["columns"]])
	# BOM: Excel otherwise reads UTF-8 accents as mojibake
	return ("\ufeff" + out.getvalue()).encode("utf-8")


def _fill_sheet(sheet, header, rows):
	from openpyxl.styles import Font

	sheet.append(header)
	for cell in sheet[1]:
		cell.font = Font(bold=True)
	for row in rows:
		sheet.append(row)
	for column in sheet.columns:
		for cell in column:
			if isinstance(cell.value, datetime.date):
				cell.number_format = "DD/MM/YYYY"
		width = max(len(_text(cell.value)) for cell in column)
		sheet.column_dimensions[column[0].column_letter].width = min(50, width + 2)
	sheet.freeze_panes = "A2"


def build_xlsx(config, docs):
	from openpyxl import Workbook

	book = Workbook()
	sheet = book.active
	sheet.title = "Documents"
	columns = config["columns"]
	_fill_sheet(sheet, [label for field, label in columns], [[_cell(doc.get(field), field) for field, label in columns] for doc in docs])

	if config.get("items"):
		child, item_columns = config["items"]
		_fill_sheet(
			book.create_sheet("Lignes"),
			["Document"] + [label for field, label in item_columns],
			[
				[doc.name] + [_cell(line.get(field), field) for field, label in item_columns]
				for doc in docs
				for line in doc.get("items") or []
			],
		)

	out = io.BytesIO()
	book.save(out)
	return out.getvalue()


def _bold_header(row):
	for cell in row.cells:
		for run in cell.paragraphs[0].runs:
			run.bold = True


def build_docx(config, docs):
	from docx import Document

	document = Document()
	for index, doc in enumerate(docs):
		if index:
			document.add_page_break()
		key_field, key_label = config["columns"][0]
		document.add_heading(f"{config['title']} {doc.get(key_field)}", level=1)
		if doc.get("company"):
			document.add_paragraph(doc.company)

		table = document.add_table(rows=0, cols=2)
		table.style = "Table Grid"
		for field, label in config["columns"]:
			cells = table.add_row().cells
			cells[0].text = label
			cells[1].text = _text(doc.get(field), field)

		if config.get("items") and doc.get("items"):
			child, item_columns = config["items"]
			document.add_heading("Lignes", level=2)
			lines = document.add_table(rows=1, cols=len(item_columns))
			lines.style = "Table Grid"
			for cell, (field, label) in zip(lines.rows[0].cells, item_columns):
				cell.text = label
			_bold_header(lines.rows[0])
			for line in doc["items"]:
				for cell, (field, label) in zip(lines.add_row().cells, item_columns):
					cell.text = _text(line.get(field), field)

	out = io.BytesIO()
	document.save(out)
	return out.getvalue()


def build_pdf(doctype, names):
	"""Each document with its default print format, one after the other in
	a single PDF. Unlike frappe's download_multi_pdf, a failing document
	raises instead of being silently left out."""
	from pypdf import PdfWriter

	writer = PdfWriter()
	for name in names:
		writer = frappe.get_print(doctype, name, as_pdf=True, output=writer)
	out = io.BytesIO()
	writer.write(out)
	return out.getvalue()
