# Copyright (c) 2026, Sitiame Capital
# License: MIT

"""Export de la balance ERPNext vers le template Excel Liasse SYSCOHADA.

Remplit les colonnes D (Débit période), E (Crédit période), F (Solde débit),
G (Solde crédit) de la feuille Balance en faisant correspondre les numéros
de comptes ERPNext aux codes 7 chiffres du template.

Le template est stocké dans les fichiers privés du site sous le nom
'liasse_syscohada_template.xlsx'.  Un fichier rempli est généré à la volée
et renvoyé en téléchargement.
"""

import io
import os
import re

import frappe
from frappe import _
from frappe.utils import get_site_path


# ── helpers ──────────────────────────────────────────────────────────────────

def _extract_account_code(account_name: str) -> str:
    """'10110 - Capital souscrit - MaPME' → '10110'"""
    m = re.match(r"^(\d+)", (account_name or "").strip())
    return m.group(1) if m else ""


def _normalize_to_7(code: str) -> str:
    """'10110' → '1011000'  (pad with trailing zeros to 7 digits, or truncate)"""
    if not code:
        return ""
    if len(code) >= 7:
        return code[:7]
    return code.ljust(7, "0")


def _get_template_path() -> str:
    private_files = get_site_path("private", "files")
    path = os.path.join(private_files, "liasse_syscohada_template.xlsx")
    if not os.path.exists(path):
        frappe.throw(
            _("Template Liasse introuvable. Uploadez d'abord le fichier "
              "'liasse_syscohada_template.xlsx' dans les fichiers privés du site."),
            frappe.ValidationError,
        )
    return path


# ── ERPNext data ──────────────────────────────────────────────────────────────

def _get_balance_data(company: str, from_date: str, to_date: str) -> dict:
    """Retourne {code_7: {debit_period, credit_period, debit_balance, credit_balance}}
    pour tous les comptes de la société sur la période."""
    gl_entries = frappe.db.sql(
        """
        SELECT
            a.account_number,
            a.account_name,
            SUM(gl.debit)  AS total_debit,
            SUM(gl.credit) AS total_credit
        FROM `tabGL Entry` gl
        JOIN `tabAccount` a ON a.name = gl.account
        WHERE
            gl.company = %(company)s
            AND gl.posting_date BETWEEN %(from_date)s AND %(to_date)s
            AND gl.is_cancelled = 0
        GROUP BY a.name
        """,
        {"company": company, "from_date": from_date, "to_date": to_date},
        as_dict=True,
    )

    # Opening balances (before from_date) to compute running balance
    opening = frappe.db.sql(
        """
        SELECT
            a.account_number,
            a.account_name,
            SUM(gl.debit)  AS open_debit,
            SUM(gl.credit) AS open_credit
        FROM `tabGL Entry` gl
        JOIN `tabAccount` a ON a.name = gl.account
        WHERE
            gl.company = %(company)s
            AND gl.posting_date < %(from_date)s
            AND gl.is_cancelled = 0
        GROUP BY a.name
        """,
        {"company": company, "from_date": from_date},
        as_dict=True,
    )

    open_map = {}
    for row in opening:
        code = _extract_account_code(row.account_number or row.account_name or "")
        key = _normalize_to_7(code)
        if key:
            open_map[key] = {
                "open_d": float(row.open_debit or 0),
                "open_c": float(row.open_credit or 0),
            }

    result = {}
    for row in gl_entries:
        code = _extract_account_code(row.account_number or row.account_name or "")
        key = _normalize_to_7(code)
        if not key:
            continue
        od = open_map.get(key, {}).get("open_d", 0.0)
        oc = open_map.get(key, {}).get("open_c", 0.0)
        td = float(row.total_debit or 0)
        tc = float(row.total_credit or 0)
        running_d = od + td
        running_c = oc + tc
        net = running_d - running_c
        result[key] = {
            "debit_period":   td,
            "credit_period":  tc,
            "debit_balance":  max(net, 0.0),
            "credit_balance": max(-net, 0.0),
        }

    return result


# ── Excel fill ────────────────────────────────────────────────────────────────

def _fill_excel(balance_data: dict, company: str, from_date: str, to_date: str) -> bytes:
    try:
        import openpyxl
    except ImportError:
        frappe.throw(_("openpyxl non installé. Lancez : pip install openpyxl"))

    wb = openpyxl.load_workbook(_get_template_path())
    ws = wb["Balance"]

    # Col B de la feuille Balance contient des formules Excel du type
    # =IFERROR(INDEX(Plan_Comptable!$B$2:$B$1456,ROW()-3),"")
    # openpyxl lit la chaîne de la formule, pas la valeur calculée.
    # On résout directement depuis Plan_Comptable : Balance row N → Plan_Comptable row (N-2).
    plan_ws = wb["Plan_Comptable"]
    # plan_codes[i] = valeur de Plan_Comptable!B(i+2), pour i >= 0
    plan_codes = []
    for pc_row in plan_ws.iter_rows(min_row=2, max_row=1456, min_col=2, max_col=2, values_only=True):
        val = pc_row[0]
        plan_codes.append(str(int(val)) if isinstance(val, (int, float)) else str(val).strip() if val else "")

    # Col indexes: A=1 B=2 C=3 D=4 E=5 F=6 G=7
    filled = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=False)):
        plan_idx = row_idx  # Balance row 4 → plan_codes[0], row 5 → plan_codes[1], …
        if plan_idx >= len(plan_codes):
            break
        raw_code = plan_codes[plan_idx]
        if not raw_code:
            continue
        code = _normalize_to_7(raw_code)
        if code in balance_data:
            data = balance_data[code]
            row[3].value = data["debit_period"]   # D
            row[4].value = data["credit_period"]  # E
            row[5].value = data["debit_balance"]  # F
            row[6].value = data["credit_balance"] # G
            filled += 1

    # Note dans PARAMETRES si la feuille existe
    if "PARAMETRES" in wb.sheetnames:
        try:
            wp = wb["PARAMETRES"]
            wp["C2"] = company
            wp["C3"] = from_date
            wp["C4"] = to_date
        except Exception:
            pass

    buf = io.BytesIO()
    wb.save(buf)
    frappe.log_error(
        title="liasse_export: balance remplie",
        message=f"{company} | {from_date}→{to_date} | {filled} comptes mappés",
    )
    return buf.getvalue()


# ── validation ───────────────────────────────────────────────────────────────

def _check_balance_equilibre(company: str, from_date: str, to_date: str):
    """Vérifie débit total = crédit total sur la période.
    Affiche un avertissement orange si l'écart dépasse 1 XOF (arrondi)."""
    row = frappe.db.sql(
        """
        SELECT
            SUM(debit)            AS total_debit,
            SUM(credit)           AS total_credit,
            ABS(SUM(debit) - SUM(credit)) AS diff
        FROM `tabGL Entry`
        WHERE company = %(company)s
          AND posting_date BETWEEN %(from_date)s AND %(to_date)s
          AND is_cancelled = 0
        """,
        {"company": company, "from_date": from_date, "to_date": to_date},
        as_dict=True,
    )
    if not row or not row[0].total_debit:
        frappe.msgprint(
            _("Aucune écriture pour cette période."),
            indicator="orange",
            title=_("Balance vide"),
        )
        return
    diff = float(row[0].diff or 0)
    if diff > 1:
        frappe.msgprint(
            _(
                "La balance n'est pas équilibrée (écart {0} XOF). "
                "Vérifiez les écritures avant d'utiliser la liasse."
            ).format(f"{diff:,.0f}"),
            indicator="orange",
            title=_("Balance déséquilibrée"),
        )


# ── public API ────────────────────────────────────────────────────────────────

@frappe.whitelist()
def export_liasse(company: str, from_date: str, to_date: str):
    """Génère le fichier Excel Liasse rempli avec la balance ERPNext et le
    renvoie en téléchargement direct (réponse binaire)."""
    roles = set(frappe.get_roles())
    if not {"Accounts Manager", "System Manager"} & roles:
        frappe.throw(_("Réservé aux comptables et administrateurs."), frappe.PermissionError)
    # Faille 2026-10-12: any PME has Accounts Manager — verify it owns this company.
    frappe.has_permission("Company", "read", company, throw=True)

    _check_balance_equilibre(company, from_date, to_date)
    balance = _get_balance_data(company, from_date, to_date)
    xlsx_bytes = _fill_excel(balance, company, from_date, to_date)

    filename = f"Liasse_{company.replace(' ', '_')}_{to_date}.xlsx"
    frappe.response["filename"] = filename
    frappe.response["filecontent"] = xlsx_bytes
    frappe.response["type"] = "binary"
