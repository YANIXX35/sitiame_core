# Copyright (c) 2026, Sitiame Capital
# License: MIT

"""Tests unitaires pour liasse_export.py — sans base de données ni réseau."""

import io

import sys

import pytest

# Ajouter get_site_path et log_error au stub frappe (absents du conftest) avant d'importer liasse_export
if "frappe.utils" in sys.modules:
    sys.modules["frappe.utils"].get_site_path = lambda *args: "/tmp/fake_site"
if "frappe" in sys.modules:
    sys.modules["frappe"].log_error = lambda *args, **kwargs: None

from sitiame_core.liasse_export import (
    _extract_account_code,
    _normalize_to_7,
    _fill_excel,
)


# ── _extract_account_code ────────────────────────────────────────────────────

def test_extract_strips_name():
    assert _extract_account_code("6052 - Electricité - PME") == "6052"


def test_extract_account_number_already_clean():
    assert _extract_account_code("40110") == "40110"


def test_extract_empty():
    assert _extract_account_code("") == ""
    assert _extract_account_code(None) == ""


def test_extract_no_leading_digit():
    assert _extract_account_code("Capital") == ""


# ── _normalize_to_7 ─────────────────────────────────────────────────────────

def test_normalize_4_digits():
    assert _normalize_to_7("6052") == "6052000"


def test_normalize_already_7():
    assert _normalize_to_7("6052000") == "6052000"


def test_normalize_longer_truncates():
    assert _normalize_to_7("40110001") == "4011000"


def test_normalize_empty():
    assert _normalize_to_7("") == ""


def test_normalize_2_digits():
    assert _normalize_to_7("60") == "6000000"


# ── _fill_excel : le bug des formules Excel (bug #1) ────────────────────────
# On crée un mini-classeur openpyxl avec :
#   - Plan_Comptable col B row 2..4 : des codes numériques
#   - Balance col B rows 4..6 : des formules (comme dans le vrai template)
# et on vérifie que _fill_excel remplit quand même les colonnes D/E/F/G.

def _make_workbook_with_formula_in_balance():
    """Simule le template réel : Balance!B a des formules, Plan_Comptable!B a les vrais codes."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Plan_Comptable : B1=titre, B2=6052, B3=4011, B4=5211
    pc = wb.create_sheet("Plan_Comptable")
    pc["B1"] = "Code"
    pc["B2"] = 6052
    pc["B3"] = 4011
    pc["B4"] = 5211

    # Balance : ligne 3 = en-tête, lignes 4..6 = formules dans col B
    bal = wb.create_sheet("Balance")
    bal["B3"] = "Compte"
    bal["D3"] = "Debit periode"
    bal["E3"] = "Credit periode"
    bal["F3"] = "Solde debit"
    bal["G3"] = "Solde credit"
    for row_i, plan_row in [(4, 2), (5, 3), (6, 4)]:
        offset = row_i - 3
        bal.cell(row_i, 2).value = f"=IFERROR(INDEX(Plan_Comptable!$B$2:$B$1456,ROW()-3),\"\")"

    # PARAMETRES
    p = wb.create_sheet("PARAMETRES")
    p["C2"] = ""

    # Sauvegarder en mémoire
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_fill_excel_resolves_formulas_from_plan_comptable(tmp_path, monkeypatch):
    """Le fichier généré doit contenir les montants même quand col B est une formule."""
    import openpyxl
    from sitiame_core import liasse_export

    template_bytes = _make_workbook_with_formula_in_balance()
    template_path = tmp_path / "liasse_syscohada_template.xlsx"
    template_path.write_bytes(template_bytes)

    monkeypatch.setattr(liasse_export, "_get_template_path", lambda: str(template_path))

    balance_data = {
        "6052000": {"debit_period": 500000, "credit_period": 0,      "debit_balance": 500000, "credit_balance": 0},
        "4011000": {"debit_period": 0,      "credit_period": 300000, "debit_balance": 0,      "credit_balance": 300000},
    }

    result = _fill_excel(balance_data, "TEST PME", "2026-01-01", "2026-12-31")

    wb_out = openpyxl.load_workbook(io.BytesIO(result))
    bal = wb_out["Balance"]

    # Row 4 → Plan_Comptable!B2 → code 6052 → 6052000 → doit être rempli
    assert bal.cell(4, 4).value == 500000, "D4 (débit période 6052) devrait être 500000"
    assert bal.cell(4, 5).value == 0,      "E4 (crédit période 6052) devrait être 0"
    assert bal.cell(4, 6).value == 500000, "F4 (solde débit 6052) devrait être 500000"
    assert bal.cell(4, 7).value == 0,      "G4 (solde crédit 6052) devrait être 0"

    # Row 5 → Plan_Comptable!B3 → code 4011 → 4011000 → doit être rempli
    assert bal.cell(5, 4).value == 0,      "D5 (débit période 4011) devrait être 0"
    assert bal.cell(5, 5).value == 300000, "E5 (crédit période 4011) devrait être 300000"
    assert bal.cell(5, 6).value == 0,      "F5 (solde débit 4011) devrait être 0"
    assert bal.cell(5, 7).value == 300000, "G5 (solde crédit 4011) devrait être 300000"

    # Row 6 → Plan_Comptable!B4 → code 5211 → 5211000 → pas dans balance_data → vide
    assert bal.cell(6, 4).value is None, "D6 (5211 absent de balance_data) devrait rester vide"


def test_fill_excel_without_plan_comptable_sheet_raises(tmp_path, monkeypatch):
    """Si la feuille Plan_Comptable est absente du template, KeyError attendu."""
    import openpyxl
    from sitiame_core import liasse_export

    wb = openpyxl.Workbook()
    wb.active.title = "Balance"
    buf = io.BytesIO()
    wb.save(buf)
    template_path = tmp_path / "liasse_syscohada_template.xlsx"
    template_path.write_bytes(buf.getvalue())

    monkeypatch.setattr(liasse_export, "_get_template_path", lambda: str(template_path))

    with pytest.raises(KeyError):
        _fill_excel({}, "TEST", "2026-01-01", "2026-12-31")
