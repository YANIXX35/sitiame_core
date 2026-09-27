# Copyright (c) 2026, Sitiame Capital
# License: MIT

"""List "Telecharger" button: document selection and the four file
formats (sitiame_core.exports)."""

import csv
import datetime
import io

import frappe
import pytest

from sitiame_core import exports

CONFIG = exports.EXPORTS["Sales Invoice"]
DOCS = [
	frappe._dict(
		name="ACC-SINV-2026-00001",
		posting_date=datetime.date(2026, 9, 17),
		customer_name="Entreprise Kouassi et Fils",
		net_total=400000.0,
		total_taxes_and_charges=0.0,
		grand_total=400000.0,
		outstanding_amount=150000.5,
		status="Unpaid",
		company="SITIAME",
		items=[
			frappe._dict(item_name="Prestation de conseil comptable", qty=1.0, rate=250000.0, amount=250000.0, income_account="7061 - S"),
			frappe._dict(item_name="Licence logiciel annuelle", qty=1.0, rate=150000.0, amount=150000.0, income_account="7061 - S"),
		],
	)
]


def test_every_export_has_a_title_a_file_and_columns():
	for doctype, config in exports.EXPORTS.items():
		assert config["title"] and config["file"], doctype
		assert config["columns"][0][0] == ("item_code" if doctype == "Item" else "name"), doctype
		assert all(label for field, label in config["columns"]), doctype


@pytest.fixture
def listed(monkeypatch):
	calls = []

	def get_list(doctype, filters=None, pluck=None, limit_page_length=None, **kwargs):
		calls.append({"doctype": doctype, "filters": filters, "limit": limit_page_length})
		return [f"DOC-{i}" for i in range(min(limit_page_length, 7))]

	monkeypatch.setattr(frappe, "get_list", get_list, raising=False)
	return calls


def test_checked_rows_win_over_the_list_filters(listed):
	exports.count_documents("Sales Invoice", "csv", names='["A", "B"]', filters='[["Sales Invoice", "status", "=", "Paid"]]')
	assert listed[0]["filters"] == [["name", "in", ["A", "B"]]]


def test_list_filters_are_used_without_checked_rows(listed):
	exports.count_documents("Sales Invoice", "csv", filters='[["Sales Invoice", "status", "=", "Paid"]]')
	assert listed[0]["filters"] == [["Sales Invoice", "status", "=", "Paid"]]


def test_count_reads_one_more_than_the_limit(listed, monkeypatch):
	monkeypatch.setitem(exports.LIMITS, "docx", 5)
	assert exports.count_documents("Sales Invoice", "docx") == {"count": 6, "limit": 5}
	assert listed[0]["limit"] == 6


def _print_settings(monkeypatch, submittable=True, cancelled=0, draft=1):
	monkeypatch.setattr(frappe, "get_meta", lambda doctype: frappe._dict(is_submittable=submittable), raising=False)
	monkeypatch.setattr(
		frappe,
		"get_single",
		lambda doctype: frappe._dict(allow_print_for_cancelled=cancelled, allow_print_for_draft=draft),
		raising=False,
	)


def test_pdf_leaves_out_what_erpnext_refuses_to_print(listed, monkeypatch):
	"""Regression: one cancelled invoice in the list made the whole PDF fail
	("Not allowed to print cancelled documents")."""
	_print_settings(monkeypatch)
	result = exports.count_documents("Sales Invoice", "pdf", filters='[["Sales Invoice", "status", "=", "Paid"]]')
	assert listed[0]["filters"] == [["Sales Invoice", "status", "=", "Paid"], ["docstatus", "not in", [2]]]
	assert "annul" in result["note"]


def test_pdf_also_leaves_out_drafts_when_printing_them_is_off(listed, monkeypatch):
	_print_settings(monkeypatch, draft=0)
	exports.count_documents("Sales Invoice", "pdf", names='["A"]')
	assert listed[0]["filters"] == [["name", "in", ["A"]], ["docstatus", "not in", [2, 0]]]


def test_pdf_of_a_non_submittable_doctype_keeps_everything(listed, monkeypatch):
	_print_settings(monkeypatch, submittable=False)
	result = exports.count_documents("Item", "pdf")
	assert listed[0]["filters"] == []
	assert "note" not in result


def test_other_formats_keep_cancelled_documents(listed, monkeypatch):
	_print_settings(monkeypatch)
	exports.count_documents("Sales Invoice", "xlsx")
	assert listed[0]["filters"] == []


def test_download_refuses_too_many_documents(listed, monkeypatch):
	monkeypatch.setitem(exports.LIMITS, "csv", 5)
	with pytest.raises(frappe.ValidationError):
		exports.download("Sales Invoice", "csv")


def test_download_refuses_an_empty_selection(monkeypatch):
	monkeypatch.setattr(frappe, "get_list", lambda *a, **k: [], raising=False)
	with pytest.raises(frappe.ValidationError):
		exports.download("Sales Invoice", "csv")


def test_unknown_doctype_or_format_is_refused():
	with pytest.raises(frappe.ValidationError):
		exports.count_documents("User", "csv")
	with pytest.raises(frappe.ValidationError):
		exports.count_documents("Sales Invoice", "exe")


def test_download_sets_a_file_response(monkeypatch):
	monkeypatch.setattr(exports, "_names", lambda *a: ["ACC-SINV-2026-00001"])
	monkeypatch.setattr(exports, "_rows", lambda doctype, config, names: DOCS)
	monkeypatch.setattr(frappe.local, "response", frappe._dict(), raising=False)
	exports.download("Sales Invoice", "csv")
	response = frappe.local.response
	assert response.type == "download"
	assert response.filename.startswith("factures_de_vente_") and response.filename.endswith(".csv")
	assert response.filecontent.startswith("\ufeff".encode("utf-8"))


def test_csv_is_excel_french_friendly():
	text = exports.build_csv(CONFIG, DOCS).decode("utf-8-sig")
	rows = list(csv.reader(io.StringIO(text), delimiter=";"))
	assert rows[0][:3] == ["N°", "Date", "Client"]
	assert rows[1][:4] == ["ACC-SINV-2026-00001", "17/09/2026", "Entreprise Kouassi et Fils", "400 000"]
	assert rows[1][6] == "150 000,50"


def test_xlsx_has_documents_and_lines_sheets():
	from openpyxl import load_workbook

	book = load_workbook(io.BytesIO(exports.build_xlsx(CONFIG, DOCS)))
	assert book.sheetnames == ["Documents", "Lignes"]
	documents = list(book["Documents"].values)
	assert documents[0][0] == "N°" and documents[1][0] == "ACC-SINV-2026-00001"
	assert documents[1][3] == 400000  # a number, not text
	lines = list(book["Lignes"].values)
	assert lines[0][0] == "Document" and len(lines) == 3
	assert lines[2][1] == "Licence logiciel annuelle"


def test_docx_has_one_section_per_document():
	from docx import Document

	document = Document(io.BytesIO(exports.build_docx(CONFIG, DOCS + DOCS)))
	headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading 1")]
	assert headings == ["Facture de vente ACC-SINV-2026-00001"] * 2
	first_table = document.tables[0]
	assert first_table.rows[2].cells[0].text == "Client" and first_table.rows[2].cells[1].text == "Entreprise Kouassi et Fils"
	lines_table = document.tables[1]
	assert lines_table.rows[0].cells[0].text == "Article" and lines_table.rows[2].cells[0].text == "Licence logiciel annuelle"


def test_booleans_read_oui_non():
	assert exports._text(1, "is_stock_item") == "Oui"
	assert exports._text(0, "is_stock_item") == "Non"
